"""Plan steps 2.8b (scenarios), 2.9 (Excel export) and the valuation-range compute (arch §5.2), for 6.3."""

import io

import pytest
from openpyxl import load_workbook


def deal(client, name="Northvale Compliance"):
    return next(d for d in client.get("/api/deals").json()["items"] if d["company"]["name"] == name)


# ---------- saved scenarios ----------


def test_save_and_list_scenarios(client):
    d = deal(client)
    body = {"deal_id": d["id"], "name": "Bid at 15x", "case": "base", "overrides": {"entry_multiple": 15.0}}
    r = client.post("/api/commands/save-scenario", json=body)
    assert r.status_code == 201, r.text
    s = r.json()
    assert (s["name"], s["case"], s["overrides"]["entry_multiple"]) == ("Bid at 15x", "base", 15.0)
    listed = client.get("/api/lbo-scenarios", params={"deal_id": d["id"]}).json()
    assert [x["id"] for x in listed["items"]] == [s["id"]]
    other = deal(client, "Veritest Group")
    assert client.get("/api/lbo-scenarios", params={"deal_id": other["id"]}).json()["total"] == 0


def test_scenario_name_is_required_and_deal_must_exist(client):
    d = deal(client)
    bad = client.post(
        "/api/commands/save-scenario", json={"deal_id": d["id"], "name": " ", "case": "base", "overrides": {}}
    )
    assert bad.status_code == 422
    missing = client.post(
        "/api/commands/save-scenario", json={"deal_id": "nope", "name": "x", "case": "base", "overrides": {}}
    )
    assert missing.status_code == 404


def test_saved_overrides_reproduce_the_model(client):
    d = deal(client)
    overrides = {"entry_multiple": 14.5, "growth_y1": 0.25}
    client.post(
        "/api/commands/save-scenario",
        json={"deal_id": d["id"], "name": "s", "case": "up", "overrides": overrides},
    )
    s = client.get("/api/lbo-scenarios", params={"deal_id": d["id"]}).json()["items"][0]
    a = client.post(
        "/api/compute/lbo", json={"deal_id": d["id"], "case": s["case"], "overrides": s["overrides"]}
    ).json()
    b = client.post(
        "/api/compute/lbo", json={"deal_id": d["id"], "case": "up", "overrides": overrides}
    ).json()
    assert a["returns"] == b["returns"]


# ---------- valuation range (football field) ----------


def test_valuation_range(client):
    d = deal(client)
    r = client.post("/api/compute/valuation-range", json={"deal_id": d["id"]})
    assert r.status_code == 200, r.text
    v = r.json()
    labels = [row["label"] for row in v["rows"]]
    assert labels == ["Trading comps", "Precedent transactions", "LBO at 20–25% IRR"]
    for row in v["rows"]:
        assert 0 < row["low"] < row["high"]
    ebitda = d["financials"]["ebitda"]
    assert v["ask"] == pytest.approx(d["ask_multiple"] * ebitda)
    comps = v["rows"][0]
    median = d["segment"]["median_multiple"]
    assert (comps["low"], comps["high"]) == (
        pytest.approx((median - 1) * ebitda),
        pytest.approx((median + 1) * ebitda),
    )
    lbo = client.post("/api/compute/lbo", json={"deal_id": d["id"]}).json()
    assert v["rows"][2]["high"] == pytest.approx(lbo["max_bid"]["enterprise_value"], rel=1e-3)


def test_precedents_use_segment_transactions(client):
    d = deal(client)
    v = client.post("/api/compute/valuation-range", json={"deal_id": d["id"]}).json()
    multiples = sorted(
        t["ev_ebitda"]
        for t in client.get("/api/transactions", params={"segment_id": d["segment"]["id"]}).json()["items"]
    )
    assert multiples, "seed should have transactions in this segment"
    lo, hi = v["rows"][1]["low"] / d["financials"]["ebitda"], v["rows"][1]["high"] / d["financials"]["ebitda"]
    assert min(multiples) <= lo <= hi <= max(multiples)


# ---------- Excel export ----------


def test_excel_export(client):
    d = deal(client)
    req = {"deal_id": d["id"], "case": "base", "overrides": {"entry_multiple": 15.0}}
    r = client.post("/api/exports/lbo", json=req)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="northvale-compliance-lbo-base.xlsx"' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    assert wb.sheetnames == ["Summary", "Assumptions", "Sources & uses", "Projections", "Sensitivity"]
    lbo = client.post("/api/compute/lbo", json=req).json()
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows() if row[0].value}
    assert summary["Gross IRR"] == pytest.approx(lbo["returns"]["irr"])
    assert summary["MOIC"] == pytest.approx(lbo["returns"]["moic"])
    assert summary["Company"] == "Northvale Compliance"
    assert wb["Projections"].max_column == 1 + 5  # label column + five years
    assert wb["Sensitivity"].max_row >= 6


def test_export_rejects_invalid_structures_like_compute(client):
    d = deal(client)
    r = client.post(
        "/api/exports/lbo",
        json={"deal_id": d["id"], "overrides": {"total_leverage": 3, "senior_leverage": 3.5}},
    )
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "invalid_assumptions"
