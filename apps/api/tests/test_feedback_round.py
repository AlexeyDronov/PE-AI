"""8.4, feedback round 1 (user notes on Overview and Screen a teaser): AUM, lifecycle, teaser documents."""

import io

import pytest

from app.services.overview import fund_lifecycle

# ---------- Overview: "add fund lifecycle and AUM" ----------


def test_overview_summary_has_aum_as_nav_plus_dry_powder(client):
    s = client.get("/api/overview-summary").json()
    assert s["aum"] == pytest.approx(s["nav"] + s["dry_powder"])


def test_overview_summary_has_the_fund_lifecycle(client):
    # Seeded fund: vintage 2019, valued at 30 Sep 2026.
    assert client.get("/api/overview-summary").json()["lifecycle"] == {
        "vintage": 2019,
        "year": 8,
        "term_years": 10,
        "phase": "Harvesting",
    }


@pytest.mark.parametrize(
    ("as_of_year", "year", "phase"),
    [
        (2019, 1, "Investment period"),
        (2023, 5, "Investment period"),
        (2024, 6, "Harvesting"),
        (2028, 10, "Harvesting"),
        (2030, 12, "Extension"),
    ],
)
def test_lifecycle_phases(as_of_year, year, phase):
    life = fund_lifecycle(vintage=2019, as_of_year=as_of_year)
    assert (life.year, life.phase, life.term_years) == (year, phase, 10)


# ---------- Screen a teaser: "upload documents too - such as pdf, pptx, docs" ----------


def minimal_pdf(text: str) -> bytes:
    """A one-page PDF with a single line of Helvetica text, built by hand (no extra test dependency)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx_bytes(slides: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for text in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = text
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def upload(client, name: str, data: bytes, content_type: str = "application/octet-stream"):
    return client.post("/api/extract-text", files={"file": (name, data, content_type)})


def test_extracts_text_from_a_pdf(client):
    r = upload(
        client, "birch.pdf", minimal_pdf("Project Birch revenue GBP 48m EBITDA GBP 9.6m"), "application/pdf"
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert "Project Birch revenue GBP 48m" in body["text"]
    assert (body["filename"], body["pages"]) == ("birch.pdf", 1)


def test_extracts_text_from_a_word_document(client):
    r = upload(client, "birch.docx", docx_bytes(["Project Birch", "Revenue £48m and EBITDA £9.6m."]))
    assert r.status_code == 200, r.text
    assert "Project Birch\n" in r.json()["text"]
    assert "Revenue £48m and EBITDA £9.6m." in r.json()["text"]


def test_extracts_text_from_a_powerpoint_deck_slide_by_slide(client):
    r = upload(client, "birch.pptx", pptx_bytes(["Project Birch overview", "FY26 revenue £48m"]))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["text"].index("Project Birch overview") < body["text"].index("FY26 revenue £48m")
    assert body["pages"] == 2


def test_plain_text_and_markdown_pass_through(client):
    for name in ("birch.txt", "birch.md"):
        r = upload(client, name, "Project Birch: revenue £48m".encode(), "text/plain")
        assert r.status_code == 200, r.text
        assert r.json()["text"] == "Project Birch: revenue £48m"


def test_rejects_unsupported_files(client):
    r = upload(client, "birch.xlsx", b"PK\x03\x04")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "unsupported_file"


def test_rejects_documents_with_no_readable_text(client):
    r = upload(client, "scan.pdf", minimal_pdf(""), "application/pdf")
    assert r.status_code == 422
    body = r.json()["error"]
    assert body["code"] == "no_text"
    assert "scan.pdf" in body["message"]


def test_rejects_corrupt_documents_cleanly(client):
    r = upload(client, "broken.docx", b"not a zip")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "unreadable_file"


def test_rejects_files_over_10mb(client):
    r = upload(client, "huge.txt", b"a" * (10 * 1024 * 1024 + 1), "text/plain")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "file_too_large"


def test_extraction_writes_nothing(client):
    before = client.get("/api/teasers").json()["total"]
    upload(client, "birch.txt", b"Project Birch", "text/plain")
    assert client.get("/api/teasers").json()["total"] == before
