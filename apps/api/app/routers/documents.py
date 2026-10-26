import io
from pathlib import PurePath

from fastapi import APIRouter, UploadFile
from pydantic import BaseModel

from app.core.errors import ExtractionFailed

router = APIRouter(prefix="/api", tags=["documents"])

MAX_BYTES = 10 * 1024 * 1024
SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md"}


class ExtractedText(BaseModel):
    filename: str
    text: str
    pages: int | None = None


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    pages = PdfReader(io.BytesIO(data)).pages
    return "\n".join(p.extract_text() or "" for p in pages), len(pages)


def _docx(data: bytes) -> tuple[str, int | None]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join(parts), None


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    slides = Presentation(io.BytesIO(data)).slides
    texts = ["\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame) for slide in slides]
    return "\n\n".join(texts), len(texts)


@router.post("/extract-text", response_model=ExtractedText)
async def extract_text(file: UploadFile) -> ExtractedText:
    name = file.filename or "upload"
    ext = PurePath(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionFailed("unsupported_file", f"{name}: use a PDF, Word, PowerPoint, .txt or .md file.")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ExtractionFailed("file_too_large", f"{name} is over the 10 MB limit.")
    try:
        if ext == ".pdf":
            text, pages = _pdf(data)
        elif ext == ".docx":
            text, pages = _docx(data)
        elif ext == ".pptx":
            text, pages = _pptx(data)
        else:
            text, pages = data.decode("utf-8"), None
    except Exception as exc:
        raise ExtractionFailed("unreadable_file", f"Could not read {name}.") from exc
    if not text.strip():
        raise ExtractionFailed("no_text", f"No readable text in {name}; it may be a scanned image.")
    return ExtractedText(filename=name, text=text, pages=pages)
