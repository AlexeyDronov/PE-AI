"""Excel workbook for an LBO run: summary, assumptions, sources and uses, projections, sensitivity."""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.compute.lbo import LboResponse
from app.services.lbo import Grid

GBP_M = '£#,##0.0"m"'
PCT = "0.0%"
MULT = '0.00"x"'
BOLD = Font(bold=True)

PROJECTION_ROWS: list[tuple[str, str, str]] = [
    ("Revenue", "revenue", GBP_M),
    ("EBITDA margin", "margin", PCT),
    ("EBITDA", "ebitda", GBP_M),
    ("D&A", "da", GBP_M),
    ("Capex", "capex", GBP_M),
    ("Change in NWC", "nwc", GBP_M),
    ("Interest", "interest", GBP_M),
    ("Tax", "tax", GBP_M),
    ("Free cash flow", "free_cash_flow", GBP_M),
    ("Debt repayment", "repayment", GBP_M),
    ("Senior debt (close)", "senior_close", GBP_M),
    ("Sub debt (close)", "sub_close", GBP_M),
    ("Cash (close)", "cash", GBP_M),
]

ASSUMPTION_ROWS: list[tuple[str, str, str]] = [
    ("Revenue (LTM)", "revenue0", GBP_M),
    ("EBITDA (LTM)", "ebitda0", GBP_M),
    ("Entry multiple", "entry_multiple", MULT),
    ("Exit multiple", "exit_multiple", MULT),
    ("Total leverage", "total_leverage", MULT),
    ("Senior leverage", "senior_leverage", MULT),
    ("Senior rate", "rate_senior", PCT),
    ("Sub rate", "rate_sub", PCT),
    ("Transaction fees", "tx_fee_pct", PCT),
    ("Financing fees", "fin_fee_pct", PCT),
    ("Minimum cash", "min_cash", GBP_M),
    ("Tax rate", "tax_rate", PCT),
    ("Margin change (bp p.a.)", "margin_change_bp", "0"),
    ("Capex (% revenue)", "capex_pct", PCT),
    ("D&A (% revenue)", "da_pct", PCT),
    ("NWC (% change in revenue)", "nwc_pct", PCT),
    ("Cash sweep", "sweep", PCT),
    ("Management rollover", "mgmt_rollover", PCT),
]

SOURCES_USES_ROWS: list[tuple[str, str]] = [
    ("Enterprise value", "enterprise_value"),
    ("Transaction fees", "transaction_fees"),
    ("Financing fees", "financing_fees"),
    ("Minimum cash", "min_cash"),
    ("Total uses", "total_uses"),
    ("Senior debt", "senior_debt"),
    ("Sub debt", "sub_debt"),
    ("Rollover equity", "rollover_equity"),
    ("Sponsor equity", "sponsor_equity"),
    ("Total sources", "total_sources"),
]


def _write(ws: Worksheet, row: int, values: list[Any], formats: list[str | None], bold: bool = False) -> None:
    for col, (value, fmt) in enumerate(zip(values, formats, strict=True), start=1):
        cell = ws.cell(row=row, column=col, value=value)
        if fmt:
            cell.number_format = fmt
        if bold:
            cell.font = BOLD


def _pairs(ws: Worksheet, header: list[str], rows: list[tuple[str, Any, str | None]]) -> None:
    _write(ws, 1, header, [None] * len(header), bold=True)
    for i, (label, value, fmt) in enumerate(rows, start=2):
        _write(ws, i, [label, value], [None, fmt])
        ws.cell(row=i, column=1).font = BOLD
    ws.freeze_panes = "A2"


def _widths(ws: Worksheet, first: float, rest: float) -> None:
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = first if col == 1 else rest


def build_lbo_workbook(company: str, model: LboResponse, grid: Grid) -> bytes:
    """Render the LBO result and sensitivity grid as an .xlsx file."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    r, su, a = model.returns, model.sources_uses, model.assumptions
    max_bid = model.max_bid
    _pairs(
        summary,
        ["Metric", "Value"],
        [
            ("Company", company, None),
            ("Case", model.case, None),
            ("Gross IRR", r.irr, PCT),
            ("MOIC", r.moic, MULT),
            ("Equity cheque", su.sponsor_equity, GBP_M),
            ("Entry EV", su.enterprise_value, GBP_M),
            ("Exit EV", r.exit_ev, GBP_M),
            ("IRR hurdle", model.hurdle, PCT),
            ("Max bid (EV)", max_bid.enterprise_value if max_bid else None, GBP_M),
            ("Max bid multiple", max_bid.multiple if max_bid else None, MULT),
        ],
    )
    _widths(summary, 24, 28)

    _pairs(
        wb.create_sheet("Assumptions"),
        ["Assumption", "Value"],
        [(label, getattr(a, key), fmt) for label, key, fmt in ASSUMPTION_ROWS]
        + [(f"Revenue growth, year {i}", g, PCT) for i, g in enumerate(a.revenue_growth, start=1)],
    )
    _pairs(
        wb.create_sheet("Sources & uses"),
        ["Item", "£m"],
        [(label, getattr(su, key), GBP_M) for label, key in SOURCES_USES_ROWS],
    )

    proj = wb.create_sheet("Projections")
    years = len(model.schedule)
    _write(proj, 1, ["£m"] + [f"Year {y.year}" for y in model.schedule], [None] * (years + 1), bold=True)
    for i, (label, key, fmt) in enumerate(PROJECTION_ROWS, start=2):
        _write(proj, i, [label] + [getattr(y, key) for y in model.schedule], [None] + [fmt] * years)
        proj.cell(row=i, column=1).font = BOLD
    proj.freeze_panes = "B2"

    sens = wb.create_sheet("Sensitivity")
    n = len(grid.exit)
    _write(sens, 1, ["Entry \\ Exit multiple"] + grid.exit, [None] + [MULT] * n, bold=True)
    for i, (entry, irrs) in enumerate(zip(grid.entry, grid.irr, strict=True), start=2):
        _write(sens, i, [entry] + irrs, [MULT] + [PCT] * n)
        sens.cell(row=i, column=1).font = BOLD
    sens.freeze_panes = "B2"

    for ws, first, rest in (
        (wb["Assumptions"], 30, 14),
        (wb["Sources & uses"], 24, 14),
        (proj, 24, 14),
        (sens, 22, 12),
    ):
        _widths(ws, first, rest)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
